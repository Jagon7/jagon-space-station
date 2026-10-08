"""
台股 ETF 每日持股來源（各投信官網公開的持股 / 申購買回清單）。

每個來源函式 fetch(code, name) 回傳 Snapshot：
    {"date": "YYYY-MM-DD",          # 持股基準日（T），不是 PCF 公告生效日（T+1）
     "holdings": [{"code", "name", "shares", "weight"}],  # 只含股票（台股 + 海外），不含期貨/現金
     "aum": float|None, "units": float|None, "nav": float|None}
抓不到就丟 SourceError。新增同一家投信的 ETF 只要在 fetch_etf.py 的 ETF_LIST 加一行。
"""
from __future__ import annotations

import html
import io
import json
import re
import time
import uuid
from datetime import date, datetime, timedelta, timezone

import requests

TPE = timezone(timedelta(hours=8))
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/128.0 Safari/537.36")


class SourceError(Exception):
    pass


_session = requests.Session()
_session.headers.update({"User-Agent": UA, "Accept": "application/json, text/html, */*",
                         "Accept-Language": "zh-TW,zh;q=0.9"})


def _req(method: str, url: str, *, params=None, json_body=None, data=None, headers=None,
         verify=True, timeout=30, tries=3) -> requests.Response:
    last = None
    for i in range(tries):
        try:
            r = _session.request(method, url, params=params, json=json_body, data=data,
                                 headers=headers, verify=verify, timeout=timeout)
            if r.status_code < 500:
                return r
            last = SourceError(f"HTTP {r.status_code} {url}")
        except requests.RequestException as e:
            last = e
        time.sleep(1.5 * (i + 1))
    raise SourceError(f"{url}: {last}")


def get_json(url, **kw):
    r = _req("GET", url, **kw)
    return _decode_json(r, url)


def post_json(url, body=None, **kw):
    r = _req("POST", url, json_body=body if body is not None else {}, **kw)
    return _decode_json(r, url)


def get_text(url, **kw) -> str:
    r = _req("GET", url, **kw)
    if r.status_code != 200:
        raise SourceError(f"HTTP {r.status_code} {url}")
    r.encoding = r.encoding if r.encoding and r.encoding.lower() != "iso-8859-1" else "utf-8"
    return r.text


def _decode_json(r: requests.Response, url: str):
    if r.status_code != 200:
        raise SourceError(f"HTTP {r.status_code} {url}")
    try:
        d = r.json()
    except ValueError:
        raise SourceError(f"非 JSON 回應 {url}")
    # 有些 API 把 JSON 再包成字串
    for _ in range(2):
        if isinstance(d, str) and d.strip()[:1] in "[{":
            d = json.loads(d)
    return d


# ── 小工具 ───────────────────────────────────────────────────
def to_num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace("%", "").strip()
    try:
        return float(s)
    except ValueError:
        return None


def to_iso(v) -> str:
    """20260924 / 2026/09/24 / 2026-09-24T.. / 115/09/24 / 1150924 / /Date(ms)/ → YYYY-MM-DD"""
    s = str(v).strip()
    m = re.search(r"/Date\((-?\d+)", s)
    if m:
        return datetime.fromtimestamp(int(m.group(1)) / 1000, TPE).strftime("%Y-%m-%d")
    m = re.match(r"^(\d{2,4})[/.-](\d{1,2})[/.-](\d{1,2})", s)
    if m:
        y = int(m.group(1))
        y = y + 1911 if y < 1911 else y
        return f"{y}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    digits = re.sub(r"\D", "", s)
    if len(digits) == 8:
        return f"{digits[:4]}-{digits[4:6]}-{digits[6:]}"
    if len(digits) == 7:
        return f"{int(digits[:3]) + 1911}-{digits[3:5]}-{digits[5:]}"
    raise SourceError(f"看不懂的日期：{v!r}")


def norm_code(raw) -> str:
    c = re.sub(r"\s+", " ", str(raw or "")).strip().upper()
    return re.sub(r"^(\d{4,6}[A-Z]?)(?:\s+(?:TT|TW|TWO)|\.TWO?)$", r"\1", c)


def is_stock(code: str) -> bool:
    if re.fullmatch(r"\d{4,6}[A-Z]?", code):
        return True
    return bool(re.fullmatch(r"[A-Z0-9.\-/]{1,12} [A-Z]{2}", code))  # 海外：'NVDA US'


def make_rows(items, code_k, name_k, shares_k, weight_k, weight_scale=1.0):
    rows = []
    for x in items or []:
        if isinstance(x, (list, tuple)):
            c, n, s, w = (x[i] if i < len(x) else None for i in (code_k, name_k, shares_k, weight_k))
        else:
            c, n, s, w = x.get(code_k), x.get(name_k), x.get(shares_k), x.get(weight_k)
        code = norm_code(c)
        if not is_stock(code):
            continue
        sh, wt = to_num(s), to_num(w)
        rows.append({"code": code, "name": re.sub(r"\s+", " ", str(n or "")).strip(),
                     "shares": None if sh is None else int(round(sh)),
                     "weight": None if wt is None else round(wt * weight_scale, 4)})
    return rows


def html_table_rows(fragment: str):
    """抓 <tr> 前 4 個 <td>：[代號, 名稱, 股數, 權重]；代號重複就停（RWD 頁面常有第二份表格）。"""
    out, seen = [], set()
    for tr in re.findall(r"<tr[^>]*>(.*?)</tr>", fragment, re.S):
        tds = [re.sub(r"<[^>]+>", "", html.unescape(t)).strip()
               for t in re.findall(r"<td[^>]*>(.*?)</td>", tr, re.S)]
        if len(tds) < 4:
            continue
        rows = make_rows([tds[:4]], 0, 1, 2, 3)
        if not rows:
            continue
        if rows[0]["code"] in seen:
            break
        seen.add(rows[0]["code"])
        out += rows
    return out


def find_key(obj, key):
    if isinstance(obj, dict):
        if key in obj:
            return obj[key]
        obj = list(obj.values())
    if isinstance(obj, list):
        for v in obj:
            if isinstance(v, (dict, list)):
                r = find_key(v, key)
                if r is not None:
                    return r
    return None


def weekdays_back(n: int, start: date | None = None) -> list[date]:
    d = start or datetime.now(TPE).date()
    out = []
    while len(out) < n:
        if d.weekday() < 5:
            out.append(d)
        d -= timedelta(days=1)
    return out


def tw_isin(code: str) -> str:
    """台灣 ETF ISIN：TW000 + 代號 + Luhn 檢查碼（字母先轉成數字）。"""
    body = "TW000" + code
    digits = "".join(str(int(ch, 36)) for ch in body)
    total = 0
    for i, ch in enumerate(reversed(digits)):
        n = int(ch) * (2 if i % 2 == 0 else 1)
        total += n // 10 + n % 10
    return body + str((10 - total % 10) % 10)


def snap(day, rows, aum=None, units=None, nav=None) -> dict:
    if not rows:
        raise SourceError("沒有持股資料")
    return {"date": to_iso(day), "holdings": rows,
            "aum": to_num(aum), "units": to_num(units), "nav": to_num(nav)}


def name_match(options: list[tuple[str, str]], name: str, code: str):
    for v, t in options:
        if t == name or code in t:
            return v
    for v, t in options:
        if name and (name in t or t in name):
            return v
    base = name.replace("主動", "")
    for v, t in options:
        if base and base in t:
            return v
    return None


_cache: dict[str, object] = {}


def cached(key, loader):
    if key not in _cache:
        _cache[key] = loader()
    return _cache[key]


# ── 公告日 T+1 → 持股基準日 T（只略過週末）──────────────────────
def previous_trading_day(announce: str) -> str:
    d = datetime.strptime(announce, "%Y-%m-%d").date() - timedelta(days=1)
    while d.weekday() >= 5:
        d -= timedelta(days=1)
    return d.isoformat()


# ════════════════════════════════════════════════════════════
# 各投信
# ════════════════════════════════════════════════════════════
def yuanta(code, name=""):
    page = f"/product/detail/{code}/ratio"
    d = get_json("https://etfapi.yuantaetfs.com/ectranslation/api/bridge", params={
        "APIType": "ETFAPI", "CompanyName": "YUANTAFUNDS", "PageName": page,
        "DeviceId": str(uuid.uuid4()), "FuncId": "PCF/Daily", "AppName": "ETF", "Device": "3",
        "Platform": "ETF", "ticker": code,
    }, headers={"Referer": "https://www.yuantaetfs.com" + page})
    if isinstance(d, dict) and "PCF" not in d and isinstance(d.get("Data"), dict):
        d = d["Data"]
    pcf = (d or {}).get("PCF") or {}
    rows = make_rows(((d or {}).get("FundWeights") or {}).get("StockWeights"), "code", "name", "qty", "weights")
    if not pcf.get("trandate"):
        raise SourceError("元大回應沒有 PCF")
    return snap(pcf["trandate"], rows, pcf.get("totalav"), pcf.get("osunit"), pcf.get("nav"))


def _p_value(page, label):
    m = re.search(r"<p>\s*" + re.escape(label) + r"[^<]*</p>\s*<p>\s*([\d,.]+)\s*</p>", page)
    return m.group(1) if m else None


def fubon(code, name=""):
    page = get_text("https://websys.fsit.com.tw/FubonETF/Trade/Assets.aspx", params={"stkId": code, "lan": "TW"})
    m = re.search(r"資料日期：\s*(\d{4}/\d{1,2}/\d{1,2})", page)
    if not m:
        raise SourceError("富邦頁面找不到資料日期")
    items = re.findall(
        r'<tr>\s*<td class="tac">([^<]+)</td>\s*<td>([^<]+)</td>\s*<td class="tar">([\d,]+)</td>'
        r'\s*<td class="tar">[\d,.-]+</td>\s*<td class="tar">([\d.]+)</td>', page)
    return snap(m.group(1), make_rows(items, 0, 1, 2, 3),
                _p_value(page, "基金淨資產"), _p_value(page, "基金在外流通單位數"), _p_value(page, "基金每單位淨值"))


CATHAY = "https://cwapi.cathaysite.com.tw/api/"


def _cathay(path, params):
    d = get_json(CATHAY + path, params=params, headers={"Referer": "https://www.cathaysite.com.tw/"})
    if not d.get("success"):
        raise SourceError(f"國泰 {path}: {d.get('returnMessage')}")
    return d.get("result")


def cathay(code, name="", start=None):
    def load():
        res = _cathay("ETF/GetETFList", {"FundType": "", "PerPageCount": 9999, "status": 1})
        lst = res if isinstance(res, list) else (res or {}).get("list") or []
        return {r.get("stockCode"): r.get("fundCode") for r in lst if r.get("stockCode")}
    fc = cached("cathay", load).get(code)
    if not fc:
        raise SourceError("國泰基金清單找不到")
    assets = _cathay("ETF/GetETFAssets", {"fundCode": fc}) or {}
    days = [to_iso(assets["preDate"])] if assets.get("preDate") and not start else []
    days += [d.isoformat() for d in weekdays_back(8, start) if d.isoformat() not in days]
    for i, day in enumerate(days):
        d = get_json(CATHAY + "ETF/GetETFDetailStockList",
                     params={"fundCode": fc, "SearchDate": day.replace("-", "/")})
        rows = make_rows(d.get("result") if d.get("success") else [], "stockCode", "stockName", "volumn", "weights")
        if rows:
            if i == 0 and assets.get("preDate") and not start:
                return snap(day, rows, assets.get("fundNav"), assets.get("fundOutstandingShares"), assets.get("fundPerNav"))
            return snap(day, rows)
    raise SourceError("國泰最近幾個工作天都查無持股")


CAPITAL = "https://www.capitalfund.com.tw/CFWeb/api/etf/"


def capital(code, name=""):
    def load():
        items = post_json(CAPITAL + "items")
        return {f.get("stockNo"): f.get("fundNo") for f in items.get("data", []) if f.get("stockNo")}
    fid = cached("capital", load).get(code)
    if not fid:
        raise SourceError("群益基金清單找不到")
    d = (post_json(CAPITAL + "buyback", {"fundId": fid}) or {}).get("data") or {}
    pcf = d.get("pcf") or {}
    if not pcf.get("date2"):
        raise SourceError("群益沒有 PCF")
    return snap(pcf["date2"], make_rows(d.get("stocks"), "stocNo", "stocName", "share", "weight"),
                pcf.get("nav"), pcf.get("totUnit"), pcf.get("pUnit"))


FHT = "https://www.fhtrust.com.tw/api/"


def fuhhwa(code, name="", start=None):
    def load():
        fl = get_json(FHT + "fundList", params={"ec001": 3})
        return {f.get("etf002"): f.get("fundID") for f in fl.get("result", []) if f.get("etf002")}
    fid = cached("fuhhwa", load).get(code)
    if not fid:
        raise SourceError("復華基金清單找不到")
    for day in weekdays_back(10, start):
        res = ((get_json(FHT + "assets", params={"fundID": fid, "qDate": day.strftime("%Y/%m/%d")}) or {})
               .get("result") or [{}])[0]
        detail = [x for x in (res.get("detail") or []) if x.get("ftype") == "股票"]
        if res.get("dDate") and detail:
            return snap(res["dDate"], make_rows(detail, "stockid", "stockname", "qshare", "prate_addaccint"),
                        res.get("pcf_FundNav"), res.get("pcf_FundQissue"), res.get("pcf_Fundpnav"))
    raise SourceError("復華最近 10 個工作天都沒有資料")


EZ = "https://www.ezmoney.com.tw/ETF/Transaction/"


def president(code, name=""):
    def load():
        page = get_text(EZ + "PCF")
        m = re.search(r'id="DataFundList"[^>]*data-content="([^"]*)"', page)
        if not m:
            raise SourceError("統一 PCF 頁找不到基金清單")
        lst = json.loads(html.unescape(m.group(1)))
        return {str(f.get("sStockNo")).strip(): f.get("sFundCode") for f in lst if f.get("sStockNo")}
    fc = cached("president", load).get(code)
    if not fc:
        raise SourceError("統一基金清單找不到")
    future = datetime.now(TPE).date() + timedelta(days=30)
    d = post_json(EZ + "GetPCF", {"fundCode": fc, "date": f"{future.year - 1911}/{future:%m/%d}", "specificDate": False},
                  headers={"Referer": EZ + "PCF", "X-Requested-With": "XMLHttpRequest"})
    assets = find_key(d, "asset") or []
    stocks = next((a.get("Details") for a in assets if a.get("AssetCode") == "ST"), None) or []
    pcf = find_key(d, "pcf") or {}
    pcf = pcf[0] if isinstance(pcf, list) and pcf else pcf
    if not isinstance(pcf, dict) or not pcf.get("TranDate"):
        raise SourceError("統一回應沒有 TranDate")
    return snap(pcf["TranDate"], make_rows(stocks, "DetailCode", "DetailName", "Share", "NavRate"),
                pcf.get("TotalNav") or pcf.get("NetAsset"), pcf.get("TotalUnit") or pcf.get("OutstandingUnits"),
                pcf.get("Nav"))


CTBC = "https://www.ctbcinvestments.com.tw/API/"


def ctbc(code, name="", start=None):
    def load():
        token = find_key(post_json(CTBC + "home/AuthToken?token=www.ctbcinvestments.com"), "token")
        data = find_key(post_json(CTBC + "etf/ETFList", params={"token": token}), "Data")
        data = data.get("Data") if isinstance(data, dict) else data
        return token, {str(x.get("ETF_ID")).strip(): x.get("FID") for x in (data or []) if x.get("ETF_ID")}
    token, mapping = cached("ctbc", load)
    fid = mapping.get(code)
    if not fid:
        raise SourceError("中信 ETF 清單找不到")
    for day in weekdays_back(10, start):
        r = post_json(CTBC + "etf/ETFHoldingWeight", {"FID": fid, "StartDate": day.strftime("%Y/%m/%d")},
                      params={"token": token})
        data = r.get("Data") if isinstance(r, dict) else None
        if not data or str(r.get("ResultCode")) == "1":
            continue
        fa = (data.get("FundAssets") or [{}])[0]
        stock = next((x.get("Data") for x in (data.get("FundAssetsDetail") or []) if x.get("Code") == "STOCK"), None)
        if fa.get("資料日期") and stock:
            return snap(fa["資料日期"], make_rows(stock, "code_", "name_", "qty_", "weights_"),
                        fa.get("基金淨資產價值") or fa.get("淨資產"), fa.get("已發行受益權單位總數"),
                        fa.get("每受益權單位淨資產價值"))
    raise SourceError("中信最近 10 個工作天都沒有資料")


def _trade_info_rows(entries):
    """野村 / 安聯 / 玉山同一套系統：Stocks 陣列或 DynamicTableData（標題「股票…」）。"""
    if entries.get("Stocks"):
        return make_rows(entries["Stocks"], "CStockCode", "CStockName", "CQuantity", "CWeightsPct")
    table = next((t for t in (entries.get("DynamicTableData") or [])
                  if str(t.get("TableTitle", "")).startswith("股票")), None)
    if not table or not table.get("Rows"):
        return []
    rows = table["Rows"]
    # 安聯欄位：[序, 代號, 名稱, 股數, 權重]；玉山：[代號, 名稱, 股數, 權重]
    off = 1 if rows and len(rows[0]) >= 5 else 0
    return make_rows(rows, off, off + 1, off + 2, off + 3)


def _trade_info(url, fund_no, headers=None, start=None):
    first = (start or datetime.now(TPE).date()) + timedelta(days=1)
    for day in weekdays_back(10, first):  # Date 要剛好是 PCF 公告日，逐日往回試
        r = post_json(url, {"Type": 1, "Keyword": "", "FundNo": fund_no, "Date": day.isoformat()}, headers=headers)
        e = find_key(r, "Entries")
        if isinstance(e, list):
            e = e[0] if e else None
        if not isinstance(e, dict):
            continue
        rows = _trade_info_rows(e)
        if rows:
            base = e.get("CNavDtStr") or e.get("CNavDt")
            if not base:
                raise SourceError("回應沒有持股基準日")
            return snap(base, rows, e.get("CAnceTotalAv"), e.get("CAnceTotalIssues"), e.get("CAnceNav"))
    raise SourceError("最近 10 個工作天都沒有資料")


def nomura(code, name="", start=None):
    return _trade_info("https://www.nomurafunds.com.tw/API/ETFAPI/api/Fund/GetFundTradeInfo", code, start=start)


ALLIANZ = "https://etf.allianzgi.com.tw"


def allianz(code, name="", start=None):
    def load():
        get_text(ALLIANZ + "/list-trade")
        t = get_json(ALLIANZ + "/webapi/api/AntiForgery/GetAntiForgeryToken")
        token = t if isinstance(t, str) else (find_key(t, "token") or find_key(t, "Token") or find_key(t, "Entries"))
        h = {"X-XSRF-TOKEN": token, "Referer": ALLIANZ + "/list-trade"}
        mapping = {}
        for ty in find_key(post_json(ALLIANZ + "/webapi/api/Category/GetFundTypeDropdownOptions", {}, headers=h), "Entries") or []:
            opts = find_key(post_json(ALLIANZ + "/webapi/api/Category/GetFundDropdownOptions",
                                      {"TypeId": ty.get("Id")}, headers=h), "Entries") or []
            for o in opts:
                if o.get("SecuritiesCode"):
                    mapping[str(o["SecuritiesCode"]).strip()] = o.get("FundNo")
        return h, mapping
    h, mapping = cached("allianz", load)
    if not mapping.get(code):
        raise SourceError("安聯基金清單找不到")
    return _trade_info(ALLIANZ + "/webapi/api/Fund/GetFundTradeInfo", mapping[code], headers=h, start=start)


ESUN = "https://www.esunam.com/ETFAPI/"


def esun(code, name="", start=None):
    def load():
        out = []
        for t in find_key(post_json(ESUN + "GetETFFundTypes"), "Entries") or []:
            for f in find_key(post_json(ESUN + "GetETFFundSelectList", {"TypeID": t.get("Id")}), "Entries") or []:
                if f.get("FundNo"):
                    out.append((str(f["FundNo"]), re.sub(r"\s+", "", f.get("FundShortName") or "")))
        return out
    fno = name_match(cached("esun", load), name, code)
    if not fno:
        raise SourceError("玉山基金清單找不到")
    return _trade_info(ESUN + "GetFundTradeInfo", fno, start=start)


KGI = "https://www.kgifund.com.tw/Fund/"


def kgi(code, name=""):
    def load():
        page = html.unescape(get_text(KGI + "RedemptionList"))
        return [(v, re.sub(r"\s+", "", t))
                for v, t in re.findall(r'<option[^>]*value="([^"]+)"[^>]*>([^<]+)</option>', page)]
    fid = name_match(cached("kgi", load), name, code)
    if not fid:
        raise SourceError(f"凱基選單找不到「{name}」")
    page = html.unescape(get_text(KGI + "Detail", params={"fundID": fid}))
    m = re.search(r"持股比重\s*(?:</div>\s*<p[^>]*>)?\s*[（(]\s*(\d{4}/\d{1,2}/\d{1,2})", page)
    if not m:
        raise SourceError("凱基找不到持股比重日期")
    # 凱基標的是公告日，換回持股基準日才能和其他家對齊
    return snap(previous_trading_day(to_iso(m.group(1))), html_table_rows(page[m.end():]))


def taishin(code, name=""):
    page = get_text(f"https://www.tsit.com.tw/ETF/Home/ETFSeriesDetail/{code}")
    m = (re.search(r'id="NAV_DATE"[^>]*value="([^"]+)"', page)
         or re.search(r"資料日期[：:\s]*(\d{4}[/-]\d{1,2}[/-]\d{1,2})", page))
    i = page.find("股數")
    if not m or i < 0:
        raise SourceError("台新找不到日期或持股表")
    seg = page[i:]
    j = seg.find("口數")  # 期貨表
    return snap(m.group(1), html_table_rows(seg[:j] if j > 0 else seg))


def sinopac(code, name=""):
    page = get_text(f"https://sitc.sinopac.com/SinopacEtfs/Etfs/SinglePcf/{code}")
    m = re.search(r"資料日期[：:\s]*(\d{4}/\d{1,2}/\d{1,2})", page)
    i = page.find("證券代碼")
    if not m or i < 0:
        raise SourceError("永豐找不到日期或持股表")
    j = page.find("證券代碼", i + 4)
    return snap(m.group(1), html_table_rows(page[i:j if j > 0 else None]))


FSITC = "https://www.fsitc.com.tw/"
FSITC_IDS = {"00728": "D90", "00408A": "183", "00994A": "182"}


def firstsec(code, name=""):
    fid = FSITC_IDS.get(code) or _fsitc_scan(code)
    if not fid:
        raise SourceError("第一金找不到基金 ID")
    d = post_json(FSITC + "WebAPI.aspx/Get_hd", {"pStrFundID": fid, "pStrDate": ""})
    raw = d.get("d") if isinstance(d, dict) else d
    if isinstance(raw, str):
        raw = json.loads(raw) if raw.strip() else []
    stocks = [x for x in (raw if isinstance(raw, list) else []) if str(x.get("group", "1")) == "1"]
    if not stocks or not stocks[0].get("sdate"):
        raise SourceError("第一金沒有持股資料")
    return snap(stocks[0]["sdate"], make_rows(stocks, "A", "B", "D", "C"))


def _fsitc_scan(code, budget=240):
    """第一金沒有基金清單 API：掃 FundDetail.aspx?ID=n 找「股票代號」。"""
    t0 = time.time()
    for i in sorted(range(1, 320), key=lambda i: abs(i - 183)):
        if time.time() - t0 > budget:
            break
        try:
            page = get_text(FSITC + "FundDetail.aspx", params={"ID": i}, tries=1, timeout=15)
        except SourceError:
            continue
        m = re.search(r"股票代號</td>\s*<td[^>]*>\s*(00\d{3,4}[A-Z]?)\s*<", page)
        if m:
            FSITC_IDS[m.group(1)] = str(i)
            if m.group(1) == code:
                return str(i)
    return None


AB = "https://webapi.alliancebernstein.com/v2/funds/tw/zh-tw/investor/"


def ab(code, name=""):
    isin = tw_isin(code)
    basket = get_json(AB + f"{isin}/basket")
    as_of = find_key(basket, "asOfDate")
    if not as_of:
        raise SourceError("聯博 basket 沒有 asOfDate")
    h = get_json(AB + f"{isin}/holdings", params={"date": str(as_of)[:10]})
    items = []
    for s in find_key(h, "domesticHoldings") or []:
        if "equity" in s.get("holdingCategory", ""):
            items += s.get("holdings") or s.get("items") or ([s] if s.get("holdingCode") else [])
    return snap(as_of, make_rows(items, "holdingCode", "holding", "holdingShares", "holdingPerc"),
                find_key(basket, "aum"), find_key(basket, "shares"), find_key(basket, "nav"))


def jpmorgan(code, name="", start=None):
    from openpyxl import load_workbook
    isin = tw_isin(code)
    for day in weekdays_back(15, start):
        try:
            r = _req("GET", "https://am.jpmorgan.com/FundsMarketingHandler/excel", params={
                "type": "holding_pcf", "cusip": isin, "country": "tw", "role": "twetf",
                "locale": "zh-TW", "date": day.isoformat()}, tries=1)
        except SourceError:
            continue
        if r.status_code != 200 or r.content[:2] != b"PK":
            continue
        ws = load_workbook(io.BytesIO(r.content), read_only=True, data_only=True).worksheets[0]
        table = [["" if c is None else str(c).strip() for c in row] for row in ws.iter_rows(values_only=True)]
        hdr = next((i for i, row in enumerate(table) if any("股票代碼" in c for c in row)), None)
        if hdr is None:
            continue
        cols = table[hdr]
        ci = {k: next((i for i, c in enumerate(cols) if k in c), None) for k in ("股票代碼", "股票名稱", "股數", "權重")}
        if None in ci.values():
            continue
        title = " ".join(" ".join(row) for row in table[:hdr])
        m = re.search(r"(\d{4}-\d{2}-\d{2})", title)
        items = []
        for row in table[hdr + 1:]:
            if len(row) <= max(ci.values()) or not row[ci["股票代碼"]]:
                break
            c = row[ci["股票代碼"]].upper()
            if re.fullmatch(r"[A-Z][A-Z.\-]{0,6}", c):  # 美股裸代號補市場別
                c += " US"
            items.append([c, row[ci["股票名稱"]], row[ci["股數"]], row[ci["權重"]]])
        rows = make_rows(items, 0, 1, 2, 3)
        if rows:
            return snap(m.group(1) if m else day.isoformat(), rows)
    raise SourceError("摩根最近 15 個工作天都沒有 Excel")


FTFT = "https://www.ftft.com.tw/official/api/"


def franklin(code, name=""):
    mapping = cached("franklin", lambda: {str(x.get("StockCode")).strip(): str(x.get("FundID"))
                                           for x in (get_json(FTFT + "etf") or []) if x.get("StockCode")})
    fid = mapping.get(code)
    if not fid:
        raise SourceError("富蘭克林基金清單找不到")
    d = get_json(FTFT + f"etf/shares/{fid}", params={"date": ""})
    if not d or not d.get("AssetDate"):
        raise SourceError("富蘭克林沒有持股資料")
    asset = datetime.strptime(str(d["AssetDate"])[:19], "%Y-%m-%dT%H:%M:%S") + timedelta(hours=8)  # UTC → 台北
    secs = [x for x in (d.get("Secs") or []) if x.get("SecuritiesType", "S") == "S"]
    return snap(asset.strftime("%Y-%m-%d"), make_rows(secs, "SecuritiesCode", "SecuritiesName", "Shares", "WeightingPercentage"),
                d.get("FundNetAssetValue"), d.get("TotalUnitsOutstanding"), d.get("NetAssetValuePerUnit"))


UOB_IDS = {"00918": "88329556"}


def uob(code, name="", start=None):
    fid = UOB_IDS.get(code)
    if not fid:
        raise SourceError("大華銀沒有基金代碼")
    for day in weekdays_back(8, start):
        try:
            # 官網憑證鏈缺中繼憑證；只讀公開資料
            r = get_json("https://www.uobam.com.tw/api/WebSite/pcf",
                         params={"fundID": fid, "pcfDate": day.strftime("%Y/%m/%d")}, verify=False, tries=2)
        except SourceError:
            continue
        stocks = [x for x in (r.get("result") or []) if x.get("kind") == "stock"]
        if r.get("etf002") == code and stocks:
            return snap(r["datadate"], make_rows(stocks, "code", "cName", "qty", "weight"),
                        r.get("totalAV"), r.get("totalIssues"), r.get("nav"))
    raise SourceError("大華銀最近 8 個工作天都沒有資料")


def union(code, name=""):
    page = html.unescape(get_text("https://www.usitc.com.tw/CustCenter/BuyBackList"))
    if f"( {code} )" not in page and f"({code})" not in page:
        raise SourceError(f"聯邦頁面預設基金不是 {code}")
    m = (re.search(r"基金資產\s*(?:<[^>]+>\s*)*資料日期[：:]\s*(\d{4}-\d{1,2}-\d{1,2})", page)
         or re.search(r"資料日期[：:]\s*(?:<[^>]+>\s*)*(\d{4}-\d{1,2}-\d{1,2})", page))
    i = page.find("股票投資比例")
    if not m or i < 0:
        raise SourceError("聯邦找不到資料日期或持股表")
    return snap(m.group(1), html_table_rows(page[i:]))


HN = "https://www.hnfunds.com.tw/WEB_API/HN_OW_PROD"


def _hn_headers(token=None):
    h = {"Client_Id": "WFPAPIPublicClient", "Accept-Language": "zh-TW",
         "X-Origin-Time": datetime.now(TPE).strftime("%Y-%m-%dT%H:%M:%S+08:00")}
    if token:
        h["Authorization"] = f"Bearer {token}"
    return h


def hnitc(code, name=""):
    token = cached("hnitc", lambda: find_key(post_json(HN + "/Auth/SysLogin", {}, headers=_hn_headers()), "access_token"))
    if not token:
        raise SourceError("華南永昌取得 token 失敗")
    d = find_key(get_json(HN + f"/ETF/FundDtl/AssetSet/{code}", headers=_hn_headers(token)), "Data") or {}
    info = find_key(get_json(HN + f"/ETF/FundDtl/{code}", headers=_hn_headers(token)), "Data") or {}
    day = info.get("NavDate") or d.get("NavDate")
    if not day:
        raise SourceError("華南永昌沒有日期")
    return snap(str(day)[:10], make_rows(d.get("StockList"), "StockNo", "StockName", "Share", "Weight", weight_scale=100),
                d.get("FundSize"), d.get("OsUnit"), d.get("Punit"))


# 可以查指定日期（含以前最近一份）的來源，用於補抓歷史
DATED_SOURCES = {"cathay", "fuhhwa", "ctbc", "nomura", "allianz", "esun", "jpmorgan", "uob"}

SOURCES = {
    "yuanta": ("元大", yuanta), "fubon": ("富邦", fubon), "cathay": ("國泰", cathay),
    "capital": ("群益", capital), "fuhhwa": ("復華", fuhhwa), "president": ("統一", president),
    "ctbc": ("中信", ctbc), "nomura": ("野村", nomura), "allianz": ("安聯", allianz),
    "kgi": ("凱基", kgi), "taishin": ("台新", taishin), "sinopac": ("永豐", sinopac),
    "firstsec": ("第一金", firstsec), "ab": ("聯博", ab), "jpmorgan": ("摩根", jpmorgan),
    "franklin": ("富蘭克林華美", franklin), "uob": ("大華銀", uob), "esun": ("玉山", esun),
    "union": ("聯邦", union), "hnitc": ("華南永昌", hnitc),
}
